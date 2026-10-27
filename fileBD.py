# ============================================================
#  archivos_a_bd.py  –  Extrae texto de archivos a SQLite
# ============================================================
#  Requisitos:
#    pip install python-docx openpyxl PyPDF2
# ============================================================

import os
import sqlite3
from pathlib import Path
from datetime import datetime

# ---------- 1. EXTRACTORES POR TIPO ----------

def extraer_txt(ruta):
    """Extrae texto de archivos .txt"""
    with open(ruta, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()

def extraer_docx(ruta):
    """Extrae texto de archivos Word (.docx)"""
    from docx import Document
    doc = Document(ruta)
    return "\n".join([p.text for p in doc.paragraphs])

def extraer_xlsx(ruta):
    """Extrae texto de archivos Excel (.xlsx)"""
    from openpyxl import load_workbook
    wb = load_workbook(ruta, data_only=True)
    lineas = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lineas.append(" | ".join(str(cell) if cell is not None else "" for cell in row))
    return "\n".join(lineas)

def extraer_pdf(ruta):
    """Extrae texto de archivos PDF"""
    from PyPDF2 import PdfReader
    reader = PdfReader(ruta)
    paginas = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            paginas.append(text)
    return "\n".join(paginas)

# ---------- 2. MAPEO DE EXTRACTORES ----------

EXTRACTORES = {
    ".txt":  extraer_txt,
    ".docx": extraer_docx,
    ".xlsx": extraer_xlsx,
    ".pdf":  extraer_pdf,
}

# ---------- 3. BASE DE DATOS ----------

def crear_bd(db_path="archivos.db"):
    """Crea la base de datos SQLite y la tabla documentos"""
    conn = sqlite3.connect(db_path)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS documentos (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            filename  TEXT NOT NULL,
            filepath  TEXT NOT NULL,
            filetype  TEXT,
            content   TEXT,
            file_size INTEGER,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit()
    return conn

# ---------- 4. PROCESAMIENTO ----------

def procesar_carpeta(carpeta, db_path="archivos.db"):
    """
    Recorre una carpeta, extrae el texto de cada archivo compatible
    y lo guarda en la base de datos SQLite.
    """
    carpeta = Path(carpeta)
    conn = crear_bd(db_path)
    cursor = conn.cursor()
    insertados = 0

    for archivo in carpeta.iterdir():
        if archivo.is_file():
            ext = archivo.suffix.lower()
            if ext in EXTRACTORES:
                try:
                    texto = EXTRACTORES[ext](archivo)
                    cursor.execute("""
                        INSERT INTO documentos
                        (filename, filepath, filetype, content, file_size)
                        VALUES (?, ?, ?, ?, ?)
                    """, (
                        archivo.name,
                        str(archivo.resolve()),
                        ext,
                        texto,
                        archivo.stat().st_size
                    ))
                    insertados += 1
                    print(f"✅ {archivo.name}")
                except Exception as e:
                    print(f"❌ Error en {archivo.name}: {e}")

    conn.commit()
    conn.close()
    print(f"\n🎉 {insertados} archivos indexados en '{db_path}'")

# ---------- 5. EJECUCIÓN ----------

if __name__ == "__main__":
    # Cambia esta ruta por tu carpeta de archivos
    CARPETA = Path.home() / "Documents"
    procesar_carpeta(CARPETA)